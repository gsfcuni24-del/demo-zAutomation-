"""Deterministic helpers shared by the agents: UIR context summaries and tabular file reading."""

import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.schemas.uir import RoutineType, Tag, UIRProject
from app.services.auditor.rule_engine import is_estop

MAX_ROWS = 200
MAX_CONTEXT_CHARS = 12_000


@dataclass(slots=True)
class TabularSource:
    filename: str
    rows: list[list[str]] = field(default_factory=list)

    def as_text(self) -> str:
        return "\n".join(" | ".join(cell for cell in row) for row in self.rows)


def read_tabular(path: Path, filename: str) -> TabularSource:
    """Read CSV / XLSX / text into rows of strings. No interpretation happens here."""
    suffix = path.suffix.lower()
    rows: list[list[str]] = []
    if suffix in (".xlsx", ".xls"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            for values in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in values]
                if any(cells):
                    rows.append(cells)
                if len(rows) >= MAX_ROWS:
                    break
        finally:
            workbook.close()
    else:
        text = path.read_bytes().decode("utf-8-sig", errors="replace")
        if suffix == ".csv":
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else None
            reader = (
                csv.reader(io.StringIO(text), dialect) if dialect else csv.reader(io.StringIO(text))
            )
            rows = [[c.strip() for c in r] for r in reader if any(c.strip() for c in r)]
        else:
            rows = [[line.strip()] for line in text.splitlines() if line.strip()]
    return TabularSource(filename=filename, rows=rows[:MAX_ROWS])


def target_program(uir: UIRProject) -> str:
    for routine in uir.routines:
        if routine.type is RoutineType.MAIN:
            return routine.program
    programs = [t.program for t in uir.tags if t.program] + [r.program for r in uir.routines]
    return programs[0] if programs else "MainProgram"


def estop_tags(uir: UIRProject, program: str) -> list[Tag]:
    return [t for t in uir.tags if is_estop(t) and t.program in (None, program)]


def uir_summary(uir: UIRProject) -> str:
    """Compact JSON view of the project for prompts (truncated to keep token usage bounded)."""
    data: dict[str, Any] = {
        "name": uir.name,
        "vendor": uir.vendor,
        "tags": [
            {
                "name": t.name,
                "type": t.data_type.value,
                "scope": t.program or "controller",
                "address": t.address,
                "description": t.description,
            }
            for t in uir.tags
        ],
        "routines": [
            {
                "id": r.id,
                "type": r.type.value,
                "language": r.language.value,
                "rungs": [r_.logic for r_ in r.rungs],
            }
            for r in uir.routines
        ],
        "screens": [
            {
                "id": s.id,
                "size": [s.width, s.height],
                "widgets": [
                    [w.type.value, w.tag_ref, w.x, w.y, w.width, w.height] for w in s.widgets
                ],
            }
            for s in uir.screens
        ],
    }
    text = json.dumps(data, separators=(",", ":"))
    return text if len(text) <= MAX_CONTEXT_CHARS else text[:MAX_CONTEXT_CHARS] + "...(truncated)"
